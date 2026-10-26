# -*- coding: utf-8 -*-
"""
src/multiformato.py
===================
Motor de Produção de Ingestão Multiformato, Descompactação Cirúrgica In-Memory (`.zip`),
Gatekeeper Anti-Lixo (>= 800 caracteres úteis + bloqueio regex de e-mails/capas)
e Materialização Nativa em PDF UTF-8 para compatibilidade integral com `inspect_pdf.py`
e `pdf_to_text.py`.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import asdict, dataclass, field
import hashlib
import io
from pathlib import Path
import re
from typing import Any
import xml.etree.ElementTree as ET
import zipfile

try:
    import pymupdf as fitz
except ImportError:
    try:
        import fitz  # type: ignore
    except ImportError:
        fitz = None  # type: ignore

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None  # type: ignore

try:
    import docx
except ImportError:
    docx = None  # type: ignore

try:
    from striprtf.striprtf import rtf_to_text
except ImportError:
    rtf_to_text = None  # type: ignore


# ─────────────────────────────────────────────────────────────────────────────
# CONSTANTES DE FILTRAGEM E GATEKEEPER DE PRODUÇÃO
# ─────────────────────────────────────────────────────────────────────────────

MIN_CHARS_UTEIS_GATEKEEPER: int = 800
MIN_BYTES_ARQUIVO: int = 256
MAX_ZIP_MEMBER_BYTES: int = 80 * 1024 * 1024  # 80 MB teto de segurança in-memory por membro

EXTENSOES_TEXTO_RICO = {".pdf", ".docx", ".odt", ".rtf"}

EXTENSOES_IGNORADAS_ZIP = {
    ".dwg", ".dxf", ".plt", ".rvt", ".ifc", ".skp",
    ".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff", ".webp", ".svg", ".ico",
    ".mp4", ".avi", ".mov", ".wmv", ".mkv", ".mp3", ".wav",
    ".xlsx", ".xls", ".xlsm", ".xlsb", ".csv", ".ods",
    ".exe", ".dll", ".bat", ".sh", ".db", ".thumbs", ".ds_store",
    ".rar", ".7z", ".tar", ".gz",
}

TERMOS_VITAIS_WHITELIST_PRIORITARIA = [
    "termo aditivo", "aditivo",
    "memorial descritivo", "memorial",
    "projeto básico", "projeto basico", "projeto executivo", "projeto",
    "parecer técnico", "parecer tecnico", "parecer de engenharia", "parecer",
    "nota técnica", "nota tecnica",
    "termo de rescisão", "termo de rescisao", "rescisão", "rescisao", "distrato",
    "paralisação", "paralisacao", "suspensão", "suspensao",
    "relatório de vistoria", "relatorio de vistoria", "vistoria", "supervisão", "supervisao",
    "laudo técnico", "laudo tecnico", "laudo de engenharia", "laudo",
    "medição", "medicao", "boletim de medição", "boletim de medicao",
    "restrição", "restricao", "inconformidade", "notificação de rescisão", "notificacao de rescisao",
    "plano de trabalho", "termo de referência", "termo de referencia",
    "estudo técnico", "estudo tecnico", "edital", "contrato", "termo de compromisso",
    "apostilamento", "ordem de serviço", "ordem de servico", "relatório", "relatorio",
    "matriz de risco", "mapa de risco", "convênio", "convenio", "minuta", "proposta",
    "planilha orçamentária", "planilha orcamentaria", "cronograma",
]

TERMOS_EMAIL_IMEDIATOS = [
    "e-mail", "email", "outlook", "enc:", "fw:", "re:",
    "comprovante de leitura", "confirmação de leitura", "confirmacao de leitura",
    "aviso de recebimento", "protocolo de envio", "recibo de entrega", "guia de recibo",
]

_EMAIL_HEADER_PATTERNS = [
    re.compile(r"(?im)^\s*(?:de|from)\s*:\s*.+$"),
    re.compile(r"(?im)^\s*(?:para|to)\s*:\s*.+$"),
    re.compile(r"(?im)^\s*(?:assunto|subject)\s*:\s*.+$"),
    re.compile(r"(?im)^\s*(?:enviado\s+em|data|sent|date)\s*:\s*.+$"),
]

_HTML_ESPELHO_VAZIO_RE = re.compile(
    r"(?is)(?:<html|<(!doctype)\s+html|saml2p:authnrequest|http\s+post\s+binding|sess[aã]o\s+expirada|acesso\s+negado)"
)

_CAPA_OU_DOU_CURTO_RE = re.compile(
    r"(?is)^(?:\s*(?:di[aá]rio\s+oficial\s+da\s+uni[aã]o|extrato\s+de\s+(?:contrato|termo\s+aditivo|conv[eê]nio)|"
    r"folha\s+de\s+rosto|capa\s+de\s+processo|termo\s+de\s+abertura\s+de\s+processo)\b)"
)


@dataclass
class DocumentoProcessado:
    """Representa um documento extraído em memória e avaliado pelo Gatekeeper Anti-Lixo."""
    nome_origem: str
    caminho_interno_zip: str | None
    formato_detectado: str
    tamanho_bytes: int
    sha256: str
    aprovado_gatekeeper: bool
    motivo_gatekeeper: str
    chars_brutos: int
    chars_uteis_corridos: int
    paragrafos_validos: int
    eh_email_bloqueado: bool = False
    eh_espelho_html: bool = False
    texto_util_preview: str = ""
    texto_completo_limpo: str = ""
    raw_bytes: bytes = field(default=b"", repr=False)
    erros: list[str] = field(default_factory=list)

    def to_summary_dict(self) -> dict[str, Any]:
        d = asdict(self)
        d.pop("texto_completo_limpo", None)
        d.pop("raw_bytes", None)
        return d


@dataclass
class ResultadoInspecaoZip:
    """Sumário da inspeção cirúrgica de um arquivo .zip in-memory."""
    nome_zip: str
    total_entradas_zip: int = 0
    entradas_ignoradas_engenharia_midia: list[str] = field(default_factory=list)
    entradas_texto_rico_extraidas: list[str] = field(default_factory=list)
    documentos_processados: list[DocumentoProcessado] = field(default_factory=list)
    erros: list[str] = field(default_factory=list)


def detectar_formato_bytes(raw_bytes: bytes, nome_hint: str = "") -> str:
    """Detecta o formato real de um payload em memória via magic bytes e estrutura de container."""
    ext_hint = Path(nome_hint.lower()).suffix if nome_hint else ""
    if not raw_bytes:
        return "VAZIO"

    if raw_bytes[:5] == b"%PDF-" or b"%PDF-" in raw_bytes[:1024]:
        return "pdf"

    if raw_bytes.lstrip()[:5].startswith(b"{\\rtf"):
        return "rtf"

    if raw_bytes[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(io.BytesIO(raw_bytes)) as zf:
                names = set(zf.namelist())
                if "word/document.xml" in names:
                    return "docx"
                if "content.xml" in names and (
                    "mimetype" in names or any(n.startswith("META-INF/") for n in names)
                ):
                    return "odt"
                if any(n.startswith("xl/") for n in names):
                    return "xlsx"
                return "zip"
        except Exception:
            return "zip_corrompido"

    amostra_inicio = raw_bytes[:1024].decode("utf-8", errors="ignore").strip().lower()
    if amostra_inicio.startswith("<!doctype html") or amostra_inicio.startswith("<html") or "<html" in amostra_inicio[:300]:
        return "html"

    if ext_hint in (".rtf", ".docx", ".odt", ".pdf", ".zip"):
        return ext_hint.lstrip(".")

    return "desconhecido"


def extrair_texto_pdf_bytes(raw_bytes: bytes) -> tuple[str, list[str]]:
    """Extrai texto nativo de um PDF em memória usando PyMuPDF (fitz) ou pypdf."""
    erros: list[str] = []
    paginas: list[str] = []

    if fitz is not None:
        try:
            with fitz.open(stream=raw_bytes, filetype="pdf") as doc:
                if doc.is_encrypted:
                    return "", ["PDF criptografado"]
                for i in range(len(doc)):
                    txt = doc.load_page(i).get_text("text") or ""
                    if txt.strip():
                        paginas.append(txt)
            return "\n\n".join(paginas), erros
        except Exception as exc:
            erros.append(f"PyMuPDF falhou: {exc}")

    if PdfReader is not None:
        try:
            reader = PdfReader(io.BytesIO(raw_bytes))
            if reader.is_encrypted:
                return "", ["PDF criptografado"]
            for page in reader.pages:
                txt = page.extract_text() or ""
                if txt.strip():
                    paginas.append(txt)
            return "\n\n".join(paginas), erros
        except Exception as exc:
            erros.append(f"pypdf falhou: {exc}")

    return "", erros or ["Nenhum leitor de PDF disponível"]


def extrair_texto_docx_bytes(raw_bytes: bytes) -> tuple[str, list[str]]:
    """Extrai texto rico de um arquivo .docx em memória usando python-docx e fallback XML."""
    erros: list[str] = []
    blocos: list[str] = []

    if docx is not None:
        try:
            document = docx.Document(io.BytesIO(raw_bytes))
            for para in document.paragraphs:
                t = para.text.strip()
                if t:
                    blocos.append(t)
            for table in document.tables:
                for row in table.rows:
                    celulas = [c.text.strip() for c in row.cells if c.text.strip()]
                    celulas_unicas: list[str] = []
                    for c in celulas:
                        if not celulas_unicas or celulas_unicas[-1] != c:
                            celulas_unicas.append(c)
                    linha_txt = " | ".join(celulas_unicas)
                    if len(linha_txt) >= 40:
                        blocos.append(linha_txt)
            return "\n\n".join(blocos), erros
        except Exception as exc:
            erros.append(f"python-docx falhou: {exc}")

    try:
        with zipfile.ZipFile(io.BytesIO(raw_bytes)) as zf:
            xml_bytes = zf.read("word/document.xml")
            root = ET.fromstring(xml_bytes)
            textos = [elem.text.strip() for elem in root.iter() if elem.text and elem.text.strip()]
            return "\n".join(textos), erros
    except Exception as exc:
        erros.append(f"Fallback XML docx falhou: {exc}")
        return "", erros


def extrair_texto_odt_bytes(raw_bytes: bytes) -> tuple[str, list[str]]:
    """Extrai texto rico de um arquivo OpenDocument Text (.odt) em memória."""
    erros: list[str] = []
    try:
        with zipfile.ZipFile(io.BytesIO(raw_bytes)) as zf:
            if "content.xml" not in zf.namelist():
                return "", ["Arquivo ODT não contém content.xml"]
            xml_bytes = zf.read("content.xml")
            root = ET.fromstring(xml_bytes)
            paragrafos: list[str] = []
            for elem in root.iter():
                tag_local = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
                if tag_local in ("p", "h"):
                    texto_p = "".join(elem.itertext()).strip()
                    if texto_p:
                        paragrafos.append(texto_p)
            if paragrafos:
                return "\n\n".join(paragrafos), erros
            texto_all = " ".join(t.strip() for t in root.itertext() if t.strip())
            return texto_all, erros
    except Exception as exc:
        erros.append(f"Extração ODT falhou: {exc}")
        return "", erros


def extrair_texto_rtf_bytes(raw_bytes: bytes) -> tuple[str, list[str]]:
    """Extrai texto de um arquivo .rtf em memória usando striprtf ou regex fallback."""
    erros: list[str] = []
    try:
        try:
            rtf_str = raw_bytes.decode("utf-8")
        except UnicodeDecodeError:
            rtf_str = raw_bytes.decode("latin-1", errors="replace")

        if rtf_to_text is not None:
            texto = rtf_to_text(rtf_str, errors="ignore")
            return texto.strip(), erros

        texto = re.sub(r"\\[a-z]+\d*\s?", " ", rtf_str)
        texto = re.sub(r"[{}]", "", texto)
        return texto.strip(), erros
    except Exception as exc:
        erros.append(f"Extração RTF falhou: {exc}")
        return "", erros


def detectar_comprovante_email(texto: str) -> tuple[bool, str]:
    """
    Detecta prints ou comprovantes de envio de e-mail procurando cabeçalhos
    (`De:`, `Para:`, `Assunto:`, `Enviado em:`).
    """
    if not texto:
        return False, ""

    topo = texto[:2500]
    matches = [pat.pattern for pat in _EMAIL_HEADER_PATTERNS if pat.search(topo)]

    inline_de = bool(re.search(r"(?i)\b(?:de|from)\s*:\s*[^\n]{3,80}", topo))
    inline_para = bool(re.search(r"(?i)\b(?:para|to)\s*:\s*[^\n]{3,80}", topo))
    inline_assunto = bool(re.search(r"(?i)\b(?:assunto|subject)\s*:\s*[^\n]{3,120}", topo))
    inline_enviado = bool(re.search(r"(?i)\b(?:enviado\s+em|sent|data)\s*:\s*[^\n]{3,80}", topo))

    score_headers = sum([inline_de, inline_para, inline_assunto, inline_enviado])

    if len(matches) >= 3 or score_headers >= 3 or (inline_de and inline_para and inline_assunto):
        return True, f"COMPROVANTE_EMAIL_DETECTADO(headers={score_headers})"

    return False, ""


def calcular_texto_util_corrido(texto_bruto: str) -> tuple[str, int, int]:
    """
    Calcula a Densidade Semântica Útil de texto corrido excluindo cabeçalhos repetitivos
    e metadados puramente numéricos/tabulares.
    """
    if not texto_bruto:
        return "", 0, 0

    texto = texto_bruto.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", " ")
    linhas_originais = [re.sub(r"[ \t]+", " ", l).strip() for l in texto.split("\n")]
    linhas_nao_vazias = [l for l in linhas_originais if l]

    if not linhas_nao_vazias:
        return "", 0, 0

    chaves = [l.casefold() for l in linhas_nao_vazias]
    contagem = Counter(chaves)
    repetidas = {k for k, v in contagem.items() if v > 2}

    linhas_sem_cabecalho = [l for l, k in zip(linhas_nao_vazias, chaves) if k not in repetidas]

    linhas_corridas: list[str] = []
    padrao_ruido_sei = re.compile(
        r"(?i)^(?:documento\s+assinado\s+eletronicamente|a\s+autenticidade\s+deste\s+documento|"
        r"c[oó]digo\s+verificador|c[oó]digo\s+crc|p[aá]gina\s+\d+\s+de\s+\d+|"
        r"processo\s+n[º°o]?\s*[\d\./-]+|sei\s+n[º°o]?\s*\d+)$"
    )

    for l in linhas_sem_cabecalho:
        if padrao_ruido_sei.search(l):
            continue
        letras = sum(c.isalpha() for c in l)
        total_len = len(l)
        if total_len < 20:
            continue
        if letras / max(total_len, 1) < 0.45:
            continue
        if l.count("|") >= 3 and (letras / max(total_len, 1) < 0.65):
            continue
        linhas_corridas.append(l)

    texto_corrido = "\n".join(linhas_corridas).strip()
    chars_uteis = len(texto_corrido)
    paragrafos_validos = sum(1 for l in linhas_corridas if len(l) >= 50 and " " in l)

    return texto_corrido, chars_uteis, paragrafos_validos


def avaliar_gatekeeper_anti_lixo(
    texto_bruto: str,
    formato: str,
    tamanho_bytes: int,
    min_chars_uteis: int = MIN_CHARS_UTEIS_GATEKEEPER,
) -> tuple[bool, str, str, int, int, bool, bool]:
    """
    Executa as travas do Gatekeeper Anti-Lixo:
    1. Tamanho mínimo e formato válido;
    2. Bloqueio de espelhos HTML vazios / SAML;
    3. Rejeição regex de comprovantes de envio de e-mail (`De:`, `Para:`, `Assunto:`, `Enviado em:`);
    4. Densidade Semântica Mínima (>= 800 caracteres úteis de texto corrido);
    5. Bloqueio de despachos de 1 parágrafo curto, capas de processo e extratos curtos de DOU.
    """
    if tamanho_bytes < MIN_BYTES_ARQUIVO:
        return False, f"ARQUIVO_MUITO_PEQUENO({tamanho_bytes}B)", "", 0, 0, False, False

    if formato in ("VAZIO", "html", "zip_corrompido", "xlsx", "desconhecido"):
        eh_html = formato == "html" or bool(_HTML_ESPELHO_VAZIO_RE.search(texto_bruto[:2000] if texto_bruto else ""))
        return False, f"FORMATO_INVALIDO_OU_ESPELHO[{formato}]", "", 0, 0, False, eh_html

    if not texto_bruto or not texto_bruto.strip():
        return False, "SEM_TEXTO_EXTRAIVEL", "", 0, 0, False, False

    if _HTML_ESPELHO_VAZIO_RE.search(texto_bruto[:2000]):
        return False, "ESPELHO_HTML_OU_SAML_VAZIO", "", 0, 0, False, True

    eh_email, motivo_email = detectar_comprovante_email(texto_bruto)
    if eh_email:
        return False, f"BLOQUEADO_REGEX_EMAIL[{motivo_email}]", "", 0, 0, True, False

    texto_corrido, chars_uteis, paragrafos_validos = calcular_texto_util_corrido(texto_bruto)

    if _CAPA_OU_DOU_CURTO_RE.search(texto_bruto[:600]) and chars_uteis < 1200:
        return False, f"CAPA_PROCESSO_OU_DOU_CURTO({chars_uteis} chars)", texto_corrido, chars_uteis, paragrafos_validos, False, False

    if chars_uteis < min_chars_uteis:
        return (
            False,
            f"BAIXA_DENSIDADE_SEMANTICA({chars_uteis} < {min_chars_uteis} chars uteis)",
            texto_corrido,
            chars_uteis,
            paragrafos_validos,
            False,
            False,
        )

    return (
        True,
        f"APROVADO_GATEKEEPER({chars_uteis} chars uteis | {paragrafos_validos} linhas densas)",
        texto_corrido,
        chars_uteis,
        paragrafos_validos,
        False,
        False,
    )


def processar_bytes_documento(
    raw_bytes: bytes,
    nome_origem: str,
    caminho_interno_zip: str | None = None,
    min_chars_uteis: int = MIN_CHARS_UTEIS_GATEKEEPER,
) -> DocumentoProcessado:
    """Processa um documento (.pdf, .docx, .odt, .rtf) em memória e aplica o Gatekeeper."""
    file_hash = hashlib.sha256(raw_bytes).hexdigest() if raw_bytes else ""
    formato = detectar_formato_bytes(raw_bytes, nome_origem)
    erros: list[str] = []
    texto_bruto = ""

    if formato == "pdf":
        texto_bruto, erros = extrair_texto_pdf_bytes(raw_bytes)
    elif formato == "docx":
        texto_bruto, erros = extrair_texto_docx_bytes(raw_bytes)
    elif formato == "odt":
        texto_bruto, erros = extrair_texto_odt_bytes(raw_bytes)
    elif formato == "rtf":
        texto_bruto, erros = extrair_texto_rtf_bytes(raw_bytes)
    elif formato == "html":
        try:
            texto_bruto = raw_bytes.decode("utf-8", errors="ignore")
        except Exception:
            texto_bruto = ""

    aprovado, motivo, texto_corrido, chars_uteis, paragrafos, eh_email, eh_html = avaliar_gatekeeper_anti_lixo(
        texto_bruto=texto_bruto,
        formato=formato,
        tamanho_bytes=len(raw_bytes),
        min_chars_uteis=min_chars_uteis,
    )

    return DocumentoProcessado(
        nome_origem=nome_origem,
        caminho_interno_zip=caminho_interno_zip,
        formato_detectado=formato,
        tamanho_bytes=len(raw_bytes),
        sha256=file_hash,
        aprovado_gatekeeper=aprovado,
        motivo_gatekeeper=motivo,
        chars_brutos=len(texto_bruto),
        chars_uteis_corridos=chars_uteis,
        paragrafos_validos=paragrafos,
        eh_email_bloqueado=eh_email,
        eh_espelho_html=eh_html,
        texto_util_preview=texto_corrido[:400] if texto_corrido else "",
        texto_completo_limpo=texto_corrido,
        raw_bytes=raw_bytes if aprovado else b"",
        erros=erros,
    )


def inspecionar_e_extrair_zip_in_memory(
    zip_bytes: bytes,
    nome_zip: str = "pacote.zip",
    min_chars_uteis: int = MIN_CHARS_UTEIS_GATEKEEPER,
    nivel_recursao: int = 0,
) -> ResultadoInspecaoZip:
    """
    Inspeção Cirúrgica de `.zip` In-Memory via `ZipFile.infolist()`:
    Ignora `.dwg`, `.jpg`, `.mp4`, `.xlsx`, etc. e extrai apenas `.pdf`, `.docx`, `.odt` e `.rtf`.
    """
    res = ResultadoInspecaoZip(nome_zip=nome_zip)

    try:
        with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
            infolist = zf.infolist()
            res.total_entradas_zip = len([i for i in infolist if not i.is_dir()])

            for info in infolist:
                if info.is_dir():
                    continue

                nome_membro = info.filename
                ext = Path(nome_membro.lower()).suffix

                if "__macosx" in nome_membro.lower() or Path(nome_membro).name.startswith("._"):
                    res.entradas_ignoradas_engenharia_midia.append(f"{nome_membro} [METADADO_OS]")
                    continue

                if ext in EXTENSOES_IGNORADAS_ZIP:
                    res.entradas_ignoradas_engenharia_midia.append(f"{nome_membro} [{ext}]")
                    continue

                if info.file_size > MAX_ZIP_MEMBER_BYTES:
                    res.entradas_ignoradas_engenharia_midia.append(
                        f"{nome_membro} [EXCEDE_LIMITE_{info.file_size}_BYTES]"
                    )
                    continue

                if ext == ".zip" and nivel_recursao < 1:
                    try:
                        sub_bytes = zf.read(info)
                        sub_res = inspecionar_e_extrair_zip_in_memory(
                            sub_bytes,
                            nome_zip=f"{nome_zip}::{nome_membro}",
                            min_chars_uteis=min_chars_uteis,
                            nivel_recursao=nivel_recursao + 1,
                        )
                        res.entradas_ignoradas_engenharia_midia.extend(sub_res.entradas_ignoradas_engenharia_midia)
                        res.entradas_texto_rico_extraidas.extend(sub_res.entradas_texto_rico_extraidas)
                        res.documentos_processados.extend(sub_res.documentos_processados)
                    except Exception as exc_sub:
                        res.erros.append(f"Erro em sub-zip {nome_membro}: {exc_sub}")
                    continue

                if ext in EXTENSOES_TEXTO_RICO:
                    try:
                        member_bytes = zf.read(info)
                        res.entradas_texto_rico_extraidas.append(nome_membro)
                        doc_proc = processar_bytes_documento(
                            raw_bytes=member_bytes,
                            nome_origem=Path(nome_membro).name,
                            caminho_interno_zip=f"{nome_zip}::{nome_membro}",
                            min_chars_uteis=min_chars_uteis,
                        )
                        res.documentos_processados.append(doc_proc)
                    except Exception as exc_mem:
                        res.erros.append(f"Erro ao ler membro {nome_membro}: {exc_mem}")
                else:
                    res.entradas_ignoradas_engenharia_midia.append(f"{nome_membro} [EXT_NAO_LISTADA:{ext}]")

    except Exception as exc:
        res.erros.append(f"Falha ao abrir arquivo ZIP '{nome_zip}': {exc}")

    return res


def ingerir_payload_multiformato(
    raw_bytes: bytes,
    nome_origem: str,
    min_chars_uteis: int = MIN_CHARS_UTEIS_GATEKEEPER,
) -> tuple[list[DocumentoProcessado], ResultadoInspecaoZip | None]:
    """
    Ponto de entrada universal do motor de ingestão:
    - Detecta automaticamente se o payload é `.zip` ou `.pdf`, `.docx`, `.odt`, `.rtf`;
    - Se for `.zip`, realiza a inspeção cirúrgica in-memory e retorna os documentos extraídos;
    - Se for documento único, processa e avalia no Gatekeeper Anti-Lixo.
    """
    formato = detectar_formato_bytes(raw_bytes, nome_origem)
    if formato == "zip":
        inspecao_zip = inspecionar_e_extrair_zip_in_memory(
            zip_bytes=raw_bytes,
            nome_zip=nome_origem,
            min_chars_uteis=min_chars_uteis,
        )
        return inspecao_zip.documentos_processados, inspecao_zip

    doc = processar_bytes_documento(
        raw_bytes=raw_bytes,
        nome_origem=nome_origem,
        caminho_interno_zip=None,
        min_chars_uteis=min_chars_uteis,
    )
    return [doc], None


def renderizar_texto_para_pdf_bytes(texto: str, titulo: str = "") -> bytes:
    """
    Converte texto autêntico extraído de `.docx`, `.odt`, `.rtf` ou de endpoints HTML/API
    (Pareceres Técnicos, Termos Aditivos, Ajustes de PT, Termos de Rescisão, Justificativas SISMOB)
    em um documento PDF nativo (%PDF-1.7) com camada de texto UTF-8 via PyMuPDF (fitz).
    """
    conteudo = f"{titulo}\n\n{texto}".strip() if titulo else texto.strip()
    if fitz is not None:
        doc = fitz.open()
        try:
            largura, altura = 595.0, 842.0  # A4
            margem = 45.0
            fontsize = 10.0
            line_height = 13.5

            linhas_formatadas: list[str] = []
            for paragrafo in conteudo.replace("\r\n", "\n").split("\n"):
                p = paragrafo.strip()
                if not p:
                    linhas_formatadas.append("")
                    continue
                palavras = p.split(" ")
                atual = ""
                for pal in palavras:
                    if len(atual) + len(pal) + 1 <= 92:
                        atual = f"{atual} {pal}".strip()
                    else:
                        if atual:
                            linhas_formatadas.append(atual)
                        atual = pal
                if atual:
                    linhas_formatadas.append(atual)

            linhas_por_pagina = int((altura - 2 * margem) // line_height)
            for idx_ini in range(0, max(len(linhas_formatadas), 1), linhas_por_pagina):
                bloco = linhas_formatadas[idx_ini : idx_ini + linhas_por_pagina]
                page = doc.new_page(width=largura, height=altura)
                y = margem
                for linha in bloco:
                    if linha:
                        page.insert_text(
                            (margem, y),
                            linha,
                            fontsize=fontsize,
                            fontname="helv",
                        )
                    y += line_height
            return doc.tobytes()
        finally:
            doc.close()

    raise RuntimeError("PyMuPDF (fitz) é necessário para renderizar texto em PDF nativo.")


def materializar_documento_aprovado_em_pdf(
    doc: DocumentoProcessado,
    caminho_destino: Path,
    titulo_cabecalho: str = "",
) -> bool:
    """
    Persiste um `DocumentoProcessado` aprovado pelo Gatekeeper como arquivo `.pdf` válido:
    - Se o formato original já for `pdf` nativo (`%PDF-`), grava `doc.raw_bytes` diretamente;
    - Se for `docx`, `odt`, `rtf` ou texto estruturado de portal/API, renderiza o texto limpo
      em um PDF nativo UTF-8 para compatibilidade integral com `inspect_pdf.py` e `pdf_to_text.py`.
    """
    if not doc.aprovado_gatekeeper:
        return False

    caminho_destino.parent.mkdir(parents=True, exist_ok=True)
    if doc.formato_detectado == "pdf" and doc.raw_bytes and doc.raw_bytes[:5] == b"%PDF-":
        caminho_destino.write_bytes(doc.raw_bytes)
        return True

    texto_alvo = doc.texto_completo_limpo
    if not texto_alvo:
        return False

    cabecalho = titulo_cabecalho or (
        f"DOCUMENTO OFICIAL CONVERTIDO DE {doc.formato_detectado.upper()} — ORIGEM: {doc.nome_origem}"
    )
    pdf_bytes = renderizar_texto_para_pdf_bytes(texto_alvo, titulo=cabecalho)
    caminho_destino.write_bytes(pdf_bytes)
    return True
